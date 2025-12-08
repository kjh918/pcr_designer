# app/routers/export.py

from __future__ import annotations

import json
from io import BytesIO

import pandas as pd
from fastapi import APIRouter, Form, HTTPException
from fastapi.responses import StreamingResponse

from openpyxl.styles import Font  # ⬅ 추가

## internal service ## 
from app.service.amplicone_normlizer import normalize_record, df_to_normalized_records



router = APIRouter(
    prefix="/export",
    tags=["export"],
)


@router.post("/excel")
async def export_amplicons_to_excel(
    kind: str = Form(...),
    data_json: str = Form(...),
):
    try:
        rows = json.loads(data_json)
        if not isinstance(rows, list):
            raise ValueError("data_json must be a list of dicts")

        df = pd.DataFrame(rows)
        print(df.columns)
        print(df)
        buffer = BytesIO()
        with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
            sheet_name = kind[:31] if kind else "Sheet1"
            df.to_excel(writer, index=False, sheet_name=sheet_name)

            # === 여기부터 스타일링 영역 ===
            wb = writer.book
            ws = writer.sheets[sheet_name]

            # 컬럼 인덱스 찾기 (1-based)
            col_map = {name: idx + 1 for idx, name in enumerate(df.columns)}

            # 예: 컬럼 이름이 이렇다고 가정
            # forward_sequence, reverse_sequence, probe_sequence, template
            red_font = Font(color="FF0000")
            blue_font = Font(color="0000FF")

            for row_idx in range(2, len(df) + 2):  # 1행은 헤더이므로 2부터
                # forward / reverse 셀 전체를 빨간 글씨로
                if "forward_sequence" in col_map:
                    cell = ws.cell(row=row_idx, column=col_map["forward_sequence"])
                    cell.font = red_font
                if "reverse_sequence" in col_map:
                    cell = ws.cell(row=row_idx, column=col_map["reverse_sequence"])
                    cell.font = red_font

                # probe 셀 전체를 파란 글씨로
                if "probe_sequence" in col_map:
                    cell = ws.cell(row=row_idx, column=col_map["probe_sequence"])
                    cell.font = blue_font

                # template 셀 전체에 색을 칠하고 싶다면 여기도 가능
                # (부분 문자열만 색을 다르게 하는 건 엑셀에서 좀 귀찮음)
                # if "template" in col_map:
                #     cell = ws.cell(row=row_idx, column=col_map["template"])
                #     cell.font = Font(color="000000")  # 일단 검정 등

        buffer.seek(0)
        filename = f"{kind}_amplicons.xlsx"

        return StreamingResponse(
            buffer,
            media_type=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"'
            },
        )
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"엑셀 export 실패: {e}")
